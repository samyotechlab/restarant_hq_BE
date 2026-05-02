import asyncio
from collections import defaultdict
import csv
import io
import math
import os
import pandas as pd
from datetime import datetime, timezone
from typing import AsyncGenerator, List
from uuid import uuid4
from bson import ObjectId
from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
import re
from models.menu_model import (
    BulkMenuUploadResponse,
    BulkMenuUploadResult,
    CSVMenuRow,
    MenuItemCreate,
    MenuItemResponse,
    MenuItemType,
    MenuItemUpdate,
    PaginatedMenuResponse,
)


# ─────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────

def _make_search_name(name: str) -> str:
    if not name:
        return ""
    name = name.lower()
    name = re.sub(r"[^\w\s]", " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def _to_response(doc: dict) -> MenuItemResponse:
    base_price   = doc.get("base_price")   or doc.get("price") or 0.0
    online_price = doc.get("online_price") or doc.get("price") or 0.0
    dietary      = doc.get("dietary")      or doc.get("type")  or MenuItemType.VEG
    return MenuItemResponse(
        id           = str(doc["_id"]),
        item_no      = doc["item_no"],
        item_name    = doc["item_name"],
        search_name  = doc.get("search_name"),
        category     = doc["category"],
        description  = doc.get("description"),
        base_price   = base_price,
        online_price = online_price,
        dietary      = dietary,
        unit_of_sale = doc.get("unit_of_sale"),
        options      = doc.get("options"),
        upsell       = doc.get("upsell"),
        available    = doc.get("available", True),
        created_at   = doc["created_at"],
        updated_at   = doc["updated_at"],
    )


def _build_insert_doc(item: MenuItemCreate, now: datetime) -> dict:
    return {
        "item_no":      f"ITEM-{uuid4().hex[:8].upper()}",
        "item_name":    item.item_name,
        "search_name":  _make_search_name(item.item_name),
        "category":     item.category,
        "description":  item.description,
        "base_price":   item.base_price,
        "online_price": item.online_price,
        "dietary":      item.dietary.value,
        "unit_of_sale": item.unit_of_sale,
        "options":      item.options,       # List[str] or None
        "upsell":       item.upsell,
        "available":    item.available,
        "created_at":   now,
        "updated_at":   now,
    }


def _build_update_fields(item: MenuItemCreate, now: datetime) -> dict:
    return {
        "item_name":    item.item_name,
        "search_name":  _make_search_name(item.item_name),
        "category":     item.category,
        "description":  item.description,
        "base_price":   item.base_price,
        "online_price": item.online_price,
        "dietary":      item.dietary.value,
        "unit_of_sale": item.unit_of_sale,
        "options":      item.options,       # List[str] or None
        "upsell":       item.upsell,
        "updated_at":   now,
    }


def _validate_object_id(item_id: str) -> None:
    if not ObjectId.is_valid(item_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"'{item_id}' is not a valid item ID.",
        )


# ─────────────────────────────────────────────────────────────────
# File Parsers
# ─────────────────────────────────────────────────────────────────

def parse_upload_file(file_bytes: bytes, filename: str) -> List[dict]:
    """Parse CSV or Excel upload → list of raw row dicts."""
    if filename.endswith((".xlsx", ".xls")):
        df = pd.read_excel(io.BytesIO(file_bytes))
    else:
        df = pd.read_csv(io.BytesIO(file_bytes))
    df = df.where(pd.notna(df), None)
    return df.to_dict(orient="records")


async def _stream_csv(file_path: str) -> AsyncGenerator[dict, None]:
    with open(file_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            yield row


async def _stream_excel(file_path: str) -> AsyncGenerator[dict, None]:
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(row)]
            continue
        if headers:
            yield dict(zip(headers, row))
    wb.close()


# ─────────────────────────────────────────────────────────────────
# Controller
# ─────────────────────────────────────────────────────────────────

class MenuController:

    @staticmethod
    async def create_item(data: MenuItemCreate, db) -> MenuItemResponse:
        now = datetime.now(timezone.utc)
        doc = _build_insert_doc(data, now)
        result = await db["menu_items"].insert_one(doc)
        doc["_id"] = result.inserted_id
        return _to_response(doc)


    @staticmethod
    async def get_item(item_id: str, db) -> MenuItemResponse:
        _validate_object_id(item_id)
        doc = await db["menu_items"].find_one({"_id": ObjectId(item_id)})
        if not doc:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Menu item not found.")
        return _to_response(doc)


    @staticmethod
    async def get_limit_fields(db):
        cursor = db["menu_items"].find(
            {"available": True},
            {"_id": 1, "item_name": 1, "search_name": 1, "category": 1,
             "online_price": 1, "dietary": 1, "options": 1, "upsell": 1, "unit_of_sale": 1},
        )
        docs = await cursor.to_list(length=None)
        categorized = defaultdict(list)
        for item in docs:
            categorized[item.get("category", "Others")].append({
                "id":           str(item["_id"]),
                "item_name":    item.get("item_name"),
                "search_name":  item.get("search_name"),
                "price":        item.get("online_price"),
                "dietary":      item.get("dietary"),
                "unit_of_sale": item.get("unit_of_sale"),
                "options":      item.get("options"),
                "upsell":       item.get("upsell"),
            })
        return dict(categorized)


    @staticmethod
    async def get_all_items_paginated(db, page: int = 1, limit: int = 10) -> PaginatedMenuResponse:
        skip  = (page - 1) * limit
        total = await db["menu_items"].count_documents({})
        items = await db["menu_items"].find().skip(skip).limit(limit).to_list(length=None)
        return PaginatedMenuResponse(
            total_results = total,
            page          = page,
            limit         = limit,
            total_pages   = math.ceil(total / limit),
            data          = [_to_response(i) for i in items],
        )


    @staticmethod
    async def get_all_items(db) -> List[MenuItemResponse]:
        items = await db["menu_items"].find().sort("created_at", -1).to_list(length=None)
        return [_to_response(i) for i in items]


    @staticmethod
    async def update_item(item_id: str, data: MenuItemUpdate, db) -> MenuItemResponse:
        _validate_object_id(item_id)
        existing = await db["menu_items"].find_one({"_id": ObjectId(item_id)})
        if not existing:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Menu item not found.")

        update_fields = data.model_dump(exclude_none=True)
        if not update_fields:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No fields provided to update.")

        if "item_name" in update_fields:
            update_fields["search_name"] = _make_search_name(update_fields["item_name"])

        update_fields["updated_at"] = datetime.now(timezone.utc)
        await db["menu_items"].update_one({"_id": ObjectId(item_id)}, {"$set": update_fields})
        updated = await db["menu_items"].find_one({"_id": ObjectId(item_id)})
        return _to_response(updated)


    @staticmethod
    async def delete_item(item_id: str, db) -> dict:
        _validate_object_id(item_id)
        result = await db["menu_items"].delete_one({"_id": ObjectId(item_id)})
        if result.deleted_count == 0:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Menu item not found.")
        return {"deleted_id": item_id}


    # ── Bulk Upload (file path — used by background/server-side upload) ──

    @staticmethod
    async def bulk_upload_from_file_path(file_path: str, db) -> dict:
        inserted = 0
        updated  = 0
        failed   = 0

        try:
            stream = _stream_csv(file_path) if file_path.endswith(".csv") else _stream_excel(file_path)
            async for row in stream:
                res       = await MenuController._process_row(row, db)
                inserted += res["inserted"]
                updated  += res["updated"]
                failed   += res["failed"]
        finally:
            if os.path.exists(file_path):
                os.remove(file_path)

        print(f"✅ Done: {inserted} inserted, {updated} updated, {failed} failed")
        return {"inserted": inserted, "updated": updated, "failed": failed}


    @staticmethod
    async def _process_row(raw: dict, db) -> dict:
        """Upsert a single raw row. Returns counts + doc."""
        values = [v for v in raw.values() if v is not None and str(v).strip() not in ("", "nan")]
        if not values:
            return {"inserted": 0, "updated": 0, "failed": 0, "doc": None}

        try:
            if "item_name" in raw:
                item = MenuItemCreate(**raw)
            else:
                item = CSVMenuRow(**raw).to_menu_item_create()

            now      = datetime.now(timezone.utc)
            existing = await db["menu_items"].find_one({
                "item_name": item.item_name,
                "category":  item.category,
            })

            if existing:
                await db["menu_items"].update_one(
                    {"_id": existing["_id"]},
                    {"$set": _build_update_fields(item, now)},
                )
                doc = await db["menu_items"].find_one({"_id": existing["_id"]})
                return {"inserted": 0, "updated": 1, "failed": 0, "doc": doc}
            else:
                doc    = _build_insert_doc(item, now)
                result = await db["menu_items"].insert_one(doc)
                doc["_id"] = result.inserted_id
                return {"inserted": 1, "updated": 0, "failed": 0, "doc": doc}

        except Exception as e:
            print(f"❌ Row failed: {e} | row={raw}")
            return {"inserted": 0, "updated": 0, "failed": 1, "doc": None}


    # ── Bulk Upload (raw rows — used by API endpoint) ──

    @staticmethod
    async def bulk_create_items(raw_rows: List[dict], db) -> BulkMenuUploadResponse:
        results: List[BulkMenuUploadResult] = []
        inserted = updated = failed = skipped = 0

        for index, raw in enumerate(raw_rows):
            item_name_hint = str(
                raw.get("ItemName") or raw.get("item_name") or f"Row {index}"
            ).strip()

            values = [v for v in raw.values() if v is not None and str(v).strip() not in ("", "nan")]
            if not values:
                skipped += 1
                continue

            try:
                if "item_name" in raw:
                    item = MenuItemCreate(**raw)
                else:
                    item = CSVMenuRow(**raw).to_menu_item_create()

                now      = datetime.now(timezone.utc)
                existing = await db["menu_items"].find_one({
                    "item_name": item.item_name,
                    "category":  item.category,
                })

                if existing:
                    await db["menu_items"].update_one(
                        {"_id": existing["_id"]},
                        {"$set": _build_update_fields(item, now)},
                    )
                    doc = await db["menu_items"].find_one({"_id": existing["_id"]})
                    results.append(BulkMenuUploadResult(
                        index=index, success=True, item_name=item.item_name, item=_to_response(doc)
                    ))
                    updated += 1
                else:
                    doc    = _build_insert_doc(item, now)
                    result = await db["menu_items"].insert_one(doc)
                    doc["_id"] = result.inserted_id
                    results.append(BulkMenuUploadResult(
                        index=index, success=True, item_name=item.item_name, item=_to_response(doc)
                    ))
                    inserted += 1

            except Exception as e:
                results.append(BulkMenuUploadResult(
                    index=index, success=False, item_name=item_name_hint, error=str(e)
                ))
                failed += 1

        print(f"✅ Bulk done: {inserted} inserted, {updated} updated, {failed} failed, {skipped} skipped")
        return BulkMenuUploadResponse(
            total    = len(raw_rows),
            inserted = inserted,
            updated  = updated,
            failed   = failed,
            skipped  = skipped,
            results  = results,
        )